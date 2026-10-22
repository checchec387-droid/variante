"""Exportation professionnelle des rapports comparatifs au format Microsoft Excel (.xlsx)."""
from __future__ import annotations

import io
from datetime import datetime
from typing import List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.core.models import (
    EarthworkStatus,
    QuantityCategory,
    UnitCostCatalog,
    Variant,
    VariantComparison,
)


def generate_excel_report(
    comparison: VariantComparison,
    baseline_variant: Variant,
    target_variant: Variant,
    catalog: Optional[UnitCostCatalog] = None,
    project_name: str = "Projet Routier",
    author: str = "Bureau d'Études VRD",
    warnings: Optional[List[str]] = None,
) -> bytes:
    """Génère un classeur Excel complet et stylisé contenant l'ensemble des résultats de comparaison."""
    wb = openpyxl.Workbook()
    # Supprimer la feuille par défaut
    wb.remove(wb.active)

    # Styles réutilisables
    font_title = Font(name="Calibri", size=16, bold=True, color="1E293B")
    font_section = Font(name="Calibri", size=12, bold=True, color="0F172A")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_normal = Font(name="Calibri", size=11)
    font_italic = Font(name="Calibri", size=10, italic=True, color="64748B")

    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_subheader = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_highlight = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    curr = comparison.currency or "EUR"

    # =========================================================================
    # FEUILLE 1 : Synthèse & Hypothèses
    # =========================================================================
    ws_syn = wb.create_sheet(title="Synthèse & Hypothèses")
    ws_syn.views.sheetView[0].showGridLines = True

    # Titre du projet
    ws_syn.cell(row=1, column=1, value="RAPPORT COMPARATIF DES VARIANTES ROUTIÈRES").font = font_title
    ws_syn.cell(row=2, column=1, value=f"Projet : {project_name} | Date : {datetime.now().strftime('%d/%m/%Y %H:%M')} | Auteur : {author}").font = font_italic

    # Section 1 : Hypothèses générales et sources
    ws_syn.cell(row=4, column=1, value="1. HYPOTHÈSES GÉNÉRALES ET SOURCES").font = font_section

    headers_hyp = ["Paramètre / Hypothèse", "Unité", f"Référence : {baseline_variant.name}", f"Cible : {target_variant.name}"]
    for col_idx, h in enumerate(headers_hyp, 1):
        c = ws_syn.cell(row=5, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center" if col_idx > 1 else "left")

    hyp_rows = [
        ("Longueur du tracé", "m", baseline_variant.metadata.total_length_m, target_variant.metadata.total_length_m),
        ("Emprise foncière", "m²", baseline_variant.metadata.right_of_way_m2, target_variant.metadata.right_of_way_m2),
        ("Coefficient de foisonnement", "-", baseline_variant.metadata.swelling_factor, target_variant.metadata.swelling_factor),
        ("Coefficient de compactage", "-", baseline_variant.metadata.compaction_factor, target_variant.metadata.compaction_factor),
        ("Source des données", "-", getattr(baseline_variant.metadata, "description", "Métré exporté") or "Métré exporté", getattr(target_variant.metadata, "description", "Métré exporté") or "Métré exporté"),
    ]

    r_idx = 6
    for param, unit, val_base, val_tgt in hyp_rows:
        ws_syn.cell(row=r_idx, column=1, value=param).font = font_normal
        ws_syn.cell(row=r_idx, column=2, value=unit).alignment = Alignment(horizontal="center")
        c3 = ws_syn.cell(row=r_idx, column=3, value=val_base)
        c4 = ws_syn.cell(row=r_idx, column=4, value=val_tgt)
        if isinstance(val_base, (int, float)):
            c3.number_format = "#,##0.00"
            c4.number_format = "#,##0.00"
            c3.alignment = Alignment(horizontal="right")
            c4.alignment = Alignment(horizontal="right")
        for col_idx in range(1, 5):
            ws_syn.cell(row=r_idx, column=col_idx).border = thin_border
        r_idx += 1

    # Section 2 : Synthèse comparative des résultats clés
    r_idx += 2
    ws_syn.cell(row=r_idx, column=1, value="2. SYNTHÈSE DES RÉSULTATS ET ÉCARTS").font = font_section
    r_idx += 1

    headers_res = ["Indicateur Clé", "Unité", f"Référence ({baseline_variant.name})", f"Cible ({target_variant.name})", "Écart Absolu (Δ)", "Écart Relatif (%)"]
    for col_idx, h in enumerate(headers_res, 1):
        c = ws_syn.cell(row=r_idx, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center" if col_idx > 1 else "left")

    r_idx += 1

    b_bal = comparison.baseline_earthwork
    t_bal = comparison.target_earthwork

    indicators_rows = [
        ("Longueur d'axe", "m", comparison.baseline_length_m, comparison.target_length_m, comparison.delta_length_m, comparison.delta_length_percent),
        ("Emprise au sol", "m²", comparison.baseline_right_of_way_m2, comparison.target_right_of_way_m2, comparison.delta_right_of_way_m2, comparison.delta_right_of_way_percent),
        ("Déblais totaux bruts", "m³", b_bal.cut_total_m3, t_bal.cut_total_m3, comparison.delta_cut_m3, ((t_bal.cut_total_m3 - b_bal.cut_total_m3) / b_bal.cut_total_m3 * 100) if b_bal.cut_total_m3 > 0 else 0.0),
        ("Remblais compactés", "m³", b_bal.fill_total_m3, t_bal.fill_total_m3, comparison.delta_fill_m3, ((t_bal.fill_total_m3 - b_bal.fill_total_m3) / b_bal.fill_total_m3 * 100) if b_bal.fill_total_m3 > 0 else 0.0),
        ("Solde net des terres", "m³", b_bal.net_balance_m3, t_bal.net_balance_m3, t_bal.net_balance_m3 - b_bal.net_balance_m3, None),
        ("Statut mouvement des terres", "-", b_bal.status.value, t_bal.status.value, "-", None),
    ]

    if comparison.baseline_total_cost is not None and comparison.target_total_cost is not None:
        indicators_rows.append((
            f"Coût Total Estimé HT ({curr})",
            curr,
            comparison.baseline_total_cost,
            comparison.target_total_cost,
            comparison.delta_abs_total_cost,
            comparison.delta_rel_total_cost_percent
        ))

    for label, unit, v_base, v_tgt, d_abs, d_rel in indicators_rows:
        ws_syn.cell(row=r_idx, column=1, value=label).font = font_bold if "Coût" in label or "Solde" in label else font_normal
        ws_syn.cell(row=r_idx, column=2, value=unit).alignment = Alignment(horizontal="center")

        c_b = ws_syn.cell(row=r_idx, column=3, value=v_base)
        c_t = ws_syn.cell(row=r_idx, column=4, value=v_tgt)
        c_da = ws_syn.cell(row=r_idx, column=5, value=d_abs)
        c_dr = ws_syn.cell(row=r_idx, column=6, value=f"{d_rel:+.2f}%" if d_rel is not None else "-")

        for c, v in [(c_b, v_base), (c_t, v_tgt), (c_da, d_abs)]:
            if isinstance(v, (int, float)):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
        c_dr.alignment = Alignment(horizontal="right")

        for col_idx in range(1, 7):
            cell = ws_syn.cell(row=r_idx, column=col_idx)
            cell.border = thin_border
            if "Coût" in label:
                cell.fill = fill_highlight

        r_idx += 1

    # Section 3 : Avertissements et observations
    r_idx += 2
    ws_syn.cell(row=r_idx, column=1, value="3. AVERTISSEMENTS ET OBSERVATIONS TECHNIQUES").font = font_section
    r_idx += 1

    all_warnings = list(warnings or [])
    for um in comparison.unit_mismatches:
        if um not in all_warnings:
            all_warnings.append(um)
    if t_bal.status == EarthworkStatus.DEFICITAIRE:
        all_warnings.append(f"La variante {target_variant.name} présente un déficit de remblai ({t_bal.net_balance_m3:,.0f} m³). Un emprunt extérieur sera nécessaire.")
    elif t_bal.status == EarthworkStatus.EXCEDENTAIRE:
        all_warnings.append(f"La variante {target_variant.name} est excédentaire (+{t_bal.net_balance_m3:,.0f} m³). Une zone de mise en dépôt définitive doit être prévue.")

    if not all_warnings:
        all_warnings.append("Aucune anomalie ou incohérence technique majeure détectée sur ce comparatif.")

    for w in all_warnings:
        ws_syn.cell(row=r_idx, column=1, value=f"• {w}").font = font_italic
        r_idx += 1

    # Ajustement largeur colonnes Synthèse
    for col in ws_syn.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_syn.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # =========================================================================
    # FEUILLE 2 : Comparatif Détaillé Poste par Poste
    # =========================================================================
    ws_det = wb.create_sheet(title="Comparatif Détaillé")
    ws_det.views.sheetView[0].showGridLines = True

    ws_det.cell(row=1, column=1, value="MÉTRÉ COMPARATIF DÉTAILLÉ PAR POSTE DE TRAVAUX").font = font_title
    ws_det.cell(row=2, column=1, value=f"Projet : {project_name} | Référence : {comparison.baseline_name} | Cible : {comparison.target_name}").font = font_italic

    headers_det = [
        "Famille", "Code", "Désignation du poste", "Unité",
        f"Quantité ({comparison.baseline_name})", f"Quantité ({comparison.target_name})",
        "Écart Quantité (Δ)", "Écart Quantité (%)",
        f"Prix Unitaire HT ({curr})",
        f"Montant ({comparison.baseline_name})", f"Montant ({comparison.target_name})",
        f"Écart Coût (Δ {curr})", "Écart Coût (%)"
    ]

    for col_idx, h in enumerate(headers_det, 1):
        c = ws_det.cell(row=4, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center" if col_idx > 3 else "left")

    row_d = 5
    for it in comparison.items:
        fill_row = fill_zebra if row_d % 2 == 0 else PatternFill(fill_type=None)

        ws_det.cell(row=row_d, column=1, value=it.category.value).font = font_normal
        ws_det.cell(row=row_d, column=2, value=it.code).font = font_normal
        ws_det.cell(row=row_d, column=3, value=it.designation).font = font_normal
        c_u = ws_det.cell(row=row_d, column=4, value=it.unit)
        c_u.alignment = Alignment(horizontal="center")

        # Quantités
        c_bq = ws_det.cell(row=row_d, column=5, value=it.baseline_qty)
        c_tq = ws_det.cell(row=row_d, column=6, value=it.target_qty)
        d_val = it.delta_abs_qty if it.delta_abs_qty is not None else ("Incompatible" if it.unit_mismatch else "-")
        c_dq = ws_det.cell(row=row_d, column=7, value=d_val)
        c_rq = ws_det.cell(row=row_d, column=8, value=f"{it.delta_rel_qty_percent:+.2f}%" if it.delta_rel_qty_percent is not None else "-")

        for c, val in [(c_bq, it.baseline_qty), (c_tq, it.target_qty), (c_dq, d_val)]:
            if isinstance(val, (int, float)):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
            else:
                c.alignment = Alignment(horizontal="center")
        c_rq.alignment = Alignment(horizontal="right")

        # Coûts
        pu_val = it.unit_price if it.unit_price is not None else "-"
        c_pu = ws_det.cell(row=row_d, column=9, value=pu_val)
        c_bc = ws_det.cell(row=row_d, column=10, value=it.baseline_cost if it.baseline_cost is not None else "-")
        c_tc = ws_det.cell(row=row_d, column=11, value=it.target_cost if it.target_cost is not None else "-")
        c_dc = ws_det.cell(row=row_d, column=12, value=it.delta_abs_cost if it.delta_abs_cost is not None else "-")
        c_rc = ws_det.cell(row=row_d, column=13, value=f"{it.delta_rel_cost_percent:+.2f}%" if it.delta_rel_cost_percent is not None else "-")

        for c, v in [(c_pu, it.unit_price), (c_bc, it.baseline_cost), (c_tc, it.target_cost), (c_dc, it.delta_abs_cost)]:
            if isinstance(v, (int, float)):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
        c_rc.alignment = Alignment(horizontal="right")

        for col_idx in range(1, 14):
            cell = ws_det.cell(row=row_d, column=col_idx)
            cell.border = thin_border
            if fill_row.fill_type:
                cell.fill = fill_row

        row_d += 1

    # Ligne de totalisation générale
    if comparison.baseline_total_cost is not None:
        ws_det.cell(row=row_d, column=1, value="TOTAL GÉNÉRAL ESTIMÉ HT").font = font_bold
        ws_det.cell(row=row_d, column=10, value=comparison.baseline_total_cost).font = font_bold
        ws_det.cell(row=row_d, column=11, value=comparison.target_total_cost).font = font_bold
        ws_det.cell(row=row_d, column=12, value=comparison.delta_abs_total_cost).font = font_bold
        ws_det.cell(row=row_d, column=13, value=f"{comparison.delta_rel_total_cost_percent:+.2f}%" if comparison.delta_rel_total_cost_percent is not None else "-").font = font_bold

        for col_idx in [10, 11, 12]:
            c = ws_det.cell(row=row_d, column=col_idx)
            c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="right")
        ws_det.cell(row=row_d, column=13).alignment = Alignment(horizontal="right")

        for col_idx in range(1, 14):
            c = ws_det.cell(row=row_d, column=col_idx)
            c.border = thin_border
            c.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    # Ajustement largeur colonnes Détail
    for col in ws_det.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_det.column_dimensions[col_letter].width = max(min(max_len + 3, 38), 12)

    # =========================================================================
    # FEUILLE 3 : Bordereau des Prix Unitaires (BPU)
    # =========================================================================
    if catalog and catalog.prices:
        ws_bpu = wb.create_sheet(title="Bordereau des Prix (BPU)")
        ws_bpu.views.sheetView[0].showGridLines = True

        ws_bpu.cell(row=1, column=1, value=f"BORDEREAU DES PRIX UNITAIRES DE RÉFÉRENCE ({catalog.currency})").font = font_title
        ws_bpu.cell(row=2, column=1, value=f"Projet : {project_name} | Date : {datetime.now().strftime('%d/%m/%Y')}").font = font_italic

        headers_bpu = ["Code", "Désignation de la prestation", "Unité", f"Prix Unitaire HT ({catalog.currency})"]
        for col_idx, h in enumerate(headers_bpu, 1):
            c = ws_bpu.cell(row=4, column=col_idx, value=h)
            c.font = font_header
            c.fill = fill_header
            c.alignment = Alignment(horizontal="center" if col_idx in [1, 3] else "left")

        r_bpu = 5
        for p in catalog.prices.values():
            ws_bpu.cell(row=r_bpu, column=1, value=p.code).font = font_normal
            ws_bpu.cell(row=r_bpu, column=2, value=p.designation).font = font_normal
            c_u = ws_bpu.cell(row=r_bpu, column=3, value=p.unit)
            c_u.alignment = Alignment(horizontal="center")
            c_pr = ws_bpu.cell(row=r_bpu, column=4, value=p.price)
            c_pr.number_format = "#,##0.00"
            c_pr.alignment = Alignment(horizontal="right")

            for col_idx in range(1, 5):
                ws_bpu.cell(row=r_bpu, column=col_idx).border = thin_border
            r_bpu += 1

        for col in ws_bpu.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_bpu.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # Sauvegarde en flux binaire
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
