"""Automated tests for Excel and PDF reporting generators."""
import io
import openpyxl
import pytest
from src.core.calculator import compare_variants
from src.reporting.excel_exporter import generate_excel_report
from src.reporting.pdf_exporter import generate_pdf_report


def test_excel_report_generation_and_integrity(variant_a_fixture, variant_b_fixture, catalog_fixture):
    """Vérifie la création du rapport Excel, ses onglets et l'exactitude des chiffres calculés."""
    comp = compare_variants(
        baseline=variant_a_fixture,
        target=variant_b_fixture,
        catalog=catalog_fixture
    )

    excel_bytes = generate_excel_report(
        comparison=comp,
        baseline_variant=variant_a_fixture,
        target_variant=variant_b_fixture,
        catalog=catalog_fixture,
        project_name="Déviation RN20",
        author="Ingénieur VRD Test",
        warnings=["Avertissement test géotechnique"]
    )

    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    # Lecture du classeur généré avec openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    sheet_names = wb.sheetnames

    # Vérification de la présence des 3 onglets
    assert "Synthèse & Hypothèses" in sheet_names
    assert "Comparatif Détaillé" in sheet_names
    assert "Bordereau des Prix (BPU)" in sheet_names

    # Vérification des chiffres dans la feuille Synthèse
    ws_syn = wb["Synthèse & Hypothèses"]
    # Vérification de métadonnées
    title_cell = ws_syn.cell(row=1, column=1).value
    assert "RAPPORT COMPARATIF" in title_cell

    # Recherche des lignes d'indicateurs
    found_length = False
    found_cost = False
    for r in range(1, 25):
        val = ws_syn.cell(row=r, column=1).value
        if val == "Longueur d'axe":
            found_length = True
            assert ws_syn.cell(row=r, column=3).value == 2500.0  # Réf
            assert ws_syn.cell(row=r, column=4).value == 2200.0  # Cible
            assert ws_syn.cell(row=r, column=5).value == -300.0  # Delta
        elif val and "Coût Total" in str(val):
            found_cost = True
            assert ws_syn.cell(row=r, column=3).value == 1872500.00  # Réf
            assert ws_syn.cell(row=r, column=4).value == 1865600.00  # Cible
            assert ws_syn.cell(row=r, column=5).value == -6900.00    # Delta

    assert found_length is True
    assert found_cost is True

    # Vérification de la feuille Comparatif Détaillé
    ws_det = wb["Comparatif Détaillé"]
    found_dm = False
    for r in range(4, 20):
        code_val = ws_det.cell(row=r, column=2).value
        if code_val == "deblai_meuble":
            found_dm = True
            assert ws_det.cell(row=r, column=5).value == 40000.0   # Base
            assert ws_det.cell(row=r, column=6).value == 30000.0   # Target
            assert ws_det.cell(row=r, column=7).value == -10000.0  # Delta Qté
            assert ws_det.cell(row=r, column=9).value == 6.50      # Prix
            assert ws_det.cell(row=r, column=10).value == 260000.0 # Coût base
            assert ws_det.cell(row=r, column=11).value == 195000.0 # Coût cible
            assert ws_det.cell(row=r, column=12).value == -65000.0 # Delta coût

    assert found_dm is True


def test_pdf_report_generation(variant_a_fixture, variant_b_fixture, catalog_fixture):
    """Vérifie la génération du fichier PDF et sa conformité binaire."""
    comp = compare_variants(
        baseline=variant_a_fixture,
        target=variant_b_fixture,
        catalog=catalog_fixture
    )

    pdf_bytes = generate_pdf_report(
        comparison=comp,
        baseline_variant=variant_a_fixture,
        target_variant=variant_b_fixture,
        catalog=catalog_fixture,
        project_name="Liaison Express Est-Ouest",
        author="Bureau d'Ingénierie des Transports",
    )

    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 2000
    # Vérification de la signature du format PDF
    assert pdf_bytes.startswith(b"%PDF-")


def test_pdf_and_excel_report_without_catalog(variant_a_fixture, variant_b_fixture):
    """Vérifie la génération des rapports lorsqu'aucun BPU n'est configuré (comparatif technique seul)."""
    comp_nocat = compare_variants(
        baseline=variant_a_fixture,
        target=variant_b_fixture,
        catalog=None
    )

    # Excel sans BPU
    excel_bytes = generate_excel_report(
        comparison=comp_nocat,
        baseline_variant=variant_a_fixture,
        target_variant=variant_b_fixture,
        catalog=None,
    )
    assert len(excel_bytes) > 0
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    # Ne doit pas avoir d'onglet BPU
    assert "Bordereau des Prix (BPU)" not in wb.sheetnames

    # PDF sans BPU
    pdf_bytes = generate_pdf_report(
        comparison=comp_nocat,
        baseline_variant=variant_a_fixture,
        target_variant=variant_b_fixture,
        catalog=None,
    )
    assert pdf_bytes.startswith(b"%PDF-")
